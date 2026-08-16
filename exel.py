from openpyxl import load_workbook
import os

os.system("cls")

path = r"C:\Users\user\Documents\exel.py.xlsx"

print(os.path.exists(path))

workbook = load_workbook(path)
sheet = workbook.active

print("successful")
total_score = 0
count = 0

results = []

for row in sheet.iter_rows(values_only=True, min_row=2):
    print(row)
    average = (row[2] + row[3]) / 2
    print(f"{row[1]}: {average}")
    results.append((row[1], average))

results.sort(key=lambda x: x[1], reverse=True)
for rank, student in enumerate(results, start=1):
    print(
        f"{student[0]:<10} Average: {student[1]:<5} Rank: {rank}")
